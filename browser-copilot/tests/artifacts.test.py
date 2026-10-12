import unittest, subprocess, json, tempfile
from pathlib import Path
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
class Artifacts(unittest.TestCase):
    def render(self,ext,spec):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        path=Path(self.tmp.name)/('report.'+ext)
        subprocess.run(['python3','server/artifacts.py',str(path)],input=json.dumps(spec),text=True,check=True)
        return path
    def test_document_content(self):
        p=self.render('docx',{'title':'Operations','sections':[{'heading':'Evidence','body':'Revenue: 42 USD'}]})
        self.assertIn('Revenue: 42 USD','\n'.join(x.text for x in Document(p).paragraphs))
    def test_workbook_formulas_and_injection(self):
        p=self.render('xlsx',{'title':'Sales','headers':['Month','Revenue'],'rows':[['June',20],['July',22],['Total',{'formula':'=SUM(B3:B4)'}],['=HYPERLINK("bad")',0]]})
        ws=load_workbook(p).active;self.assertEqual(ws['B5'].value,'=SUM(B3:B4)');self.assertEqual(ws['A6'].data_type,'s');self.assertEqual(ws.freeze_panes,'A3')
    def test_slides_and_sources(self):
        p=self.render('pptx',{'title':'Findings','sections':[{'heading':'Result','body':'42 USD'}],'sources':['https://example.com']})
        self.assertEqual(len(Presentation(p).slides),3)
if __name__=='__main__':unittest.main()
