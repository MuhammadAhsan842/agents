"""Deterministic Office rendering. Receives structured JSON, never executable code."""
import json, sys, re
from pathlib import Path
spec=json.load(sys.stdin)
path=Path(sys.argv[1]); title=str(spec.get('title','Report'))[:200]
sections=spec.get('sections',[])[:20]
if path.suffix=='.docx':
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    doc=Document(); sec=doc.sections[0]; sec.top_margin=sec.bottom_margin=Inches(.8)
    doc.styles['Normal'].font.name='Calibri'; doc.styles['Normal'].font.size=Pt(11)
    doc.add_heading(title,0)
    for section in sections:
        doc.add_heading(str(section.get('heading',''))[:160],1)
        doc.add_paragraph(str(section.get('body',''))[:6000])
    for source in spec.get('sources',[])[:30]: doc.add_paragraph(str(source),style='Caption')
    doc.save(path)
elif path.suffix=='.xlsx':
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    def literal(value):
        text=str(value);return "'"+text if text.startswith(('=','+','-','@')) else text
    wb=Workbook();ws=wb.active;ws.title='Summary';ws.append([literal(title)]);headers=spec.get('headers',[])[:20];ws.append([literal(x) for x in headers])
    for row in spec.get('rows',[])[:500]:
        values=[]
        for value in row[:20]:
            if isinstance(value,dict) and 'formula' in value:
                formula=str(value['formula']).upper()
                # A deliberately narrow arithmetic/SUM formula grammar; no external data functions.
                if not re.fullmatch(r'=(?:SUM\([A-Z]{1,2}[1-9][0-9]*:[A-Z]{1,2}[1-9][0-9]*\)|[A-Z]{1,2}[1-9][0-9]*(?:[+*/-](?:[A-Z]{1,2}[1-9][0-9]*|[0-9]+(?:\.[0-9]+)?))+)',formula): raise ValueError('Unsupported formula')
                values.append(formula)
            elif isinstance(value,(int,float)): values.append(value)
            else:
                text=str(value)[:2000];values.append("'"+text if text.startswith(('=','+','-','@')) else text)
        ws.append(values)
    for row in ws.iter_rows(min_row=1,max_row=2):
        for cell in row:cell.fill=PatternFill('solid',fgColor='183734');cell.font=Font(color='FFFFFF',bold=True)
    for i in range(1,max(len(headers),1)+1):ws.column_dimensions[get_column_letter(i)].width=24
    ws.freeze_panes='A3';ws.auto_filter.ref=f'A2:{get_column_letter(max(len(headers),1))}{max(ws.max_row,2)}'
    if spec.get('sources'):
        sources=wb.create_sheet('Sources');sources.append(['Source']);
        for source in spec['sources'][:30]:sources.append([literal(source)])
        sources.column_dimensions['A'].width=90
    wb.save(path)
elif path.suffix=='.pptx':
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    entries=[{'heading':title,'body':str(spec.get('subtitle','Prepared with Atlas'))}]+sections
    if spec.get('sources'):entries.append({'heading':'Sources','body':'\n'.join(str(s) for s in spec['sources'][:8])})
    for i,section in enumerate(entries):
        slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string('F6F8F5')
        def box(x,y,w,h,text,size,color):
            shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=shape.text_frame;tf.word_wrap=True
            for j,line in enumerate(text.split('\n')):
                p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.text=line;p.font.size=Pt(size);p.font.color.rgb=RGBColor.from_string(color)
        box(.8,.55,11,.4,'ATLAS / RESEARCH BRIEF',12,'26725C');box(.8,1.3,11.7,1.3,str(section.get('heading',''))[:100],34,'183734');box(.8,2.9,11.4,3.5,str(section.get('body',''))[:700],22,'425D57');box(.8,6.8,11,.3,str(i+1).zfill(2),11,'70827E')
    prs.save(path)
else: raise ValueError('Unsupported artifact format')
